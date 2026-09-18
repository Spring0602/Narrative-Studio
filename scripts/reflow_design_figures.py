"""Reflow selected report diagrams; preserve surrounding content."""
from pathlib import Path
from copy import deepcopy
from xml.sax.saxutils import escape
import math, shutil
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/figures/system-design'
QA=ROOT/'.tmp/design-reflow'
DOC=ROOT/'docs/需求分析.docx'
INK='#294660'; BLUE='#DFEDF8'; GREEN='#E8F1E1'; GOLD='#FFF1CC'; PURPLE='#EFE8F7'

class Canvas:
    def __init__(self,name,h):
        self.name=name; self.h=h
        self.im=Image.new('RGB',(1600,h),'white'); self.d=ImageDraw.Draw(self.im)
        self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="{h}"><rect width="100%" height="100%" fill="white"/>']
    def text(self,xy,text,size=30,center=False):
        x,y=xy; font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',size)
        width=self.d.textlength(text,font=font); left=x-width/2 if center else x
        assert left>=0 and left+width<=1600,(self.name,text,left,width)
        self.d.text((left,y),text,font=font,fill=INK,anchor='lt')
        anchor='middle' if center else 'start'
        self.svg.append(f'<text x="{x}" y="{y}" font-family="Microsoft YaHei,sans-serif" font-size="{size}" fill="{INK}" dominant-baseline="text-before-edge" text-anchor="{anchor}">{escape(text)}</text>')
    def line(self,ps,dash=False,arrow=False,hollow=False):
        for a,b in zip(ps,ps[1:]):
            length=math.dist(a,b)
            if dash:
                for s in range(0,int(length),18):
                    e=min(s+10,length)
                    self.d.line([tuple(a[k]+(b[k]-a[k])*v/length for k in (0,1)) for v in (s,e)],fill=INK,width=3)
            else: self.d.line((a,b),fill=INK,width=3)
        coords=' '.join(f'{x},{y}' for x,y in ps); ds=' stroke-dasharray="10 8"' if dash else ''
        self.svg.append(f'<polyline points="{coords}" fill="none" stroke="{INK}" stroke-width="3"{ds}/>')
        if arrow or hollow:
            a,b=ps[-2:]; t=math.atan2(b[1]-a[1],b[0]-a[0]); l=19 if hollow else 14
            tri=[b,(b[0]-l*math.cos(t)+l*.5*math.sin(t),b[1]-l*math.sin(t)-l*.5*math.cos(t)),(b[0]-l*math.cos(t)-l*.5*math.sin(t),b[1]-l*math.sin(t)+l*.5*math.cos(t))]
            fill='white' if hollow else INK; self.d.polygon(tri,fill=fill,outline=INK,width=3)
            coords=' '.join(f'{x},{y}' for x,y in tri)
            self.svg.append(f'<polygon points="{coords}" fill="{fill}" stroke="{INK}" stroke-width="3"/>')
    def rect(self,xy,fill='white'):
        x,y,r,b=xy; self.d.rectangle(xy,fill=fill,outline=INK,width=3)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{r-x}" height="{b-y}" fill="{fill}" stroke="{INK}" stroke-width="3"/>')
    def box(self,xy,label,fill=BLUE,size=32):
        self.rect(xy,fill); x,y,r,b=xy; lines=label.split('\n'); top=(y+b-len(lines)*(size+10))/2+5
        for i,s in enumerate(lines): self.text(((x+r)/2,top+i*(size+10)),s,size,True)
    def entity(self,x,y,name,rows,fill=BLUE):
        h=83+len(rows)*43; self.rect((x,y,x+440,y+h)); self.rect((x,y,x+440,y+64),fill)
        self.text((x+220,y+17),name,28,True)
        for i,row in enumerate(rows): self.text((x+16,y+78+i*43),row,27)
    def rel(self,ps,a='1',b='0..*',label=None):
        self.line(ps); x,y=ps[0]; nx,ny=ps[1]; ex,ey=ps[-1]; px,py=ps[-2]
        if len(ps)==2 and y==ey and abs(ex-x)<180:
            self.text((x+(10 if ex>x else -26),y-36),a,25)
            self.text((ex+(-64 if ex>x else 10),ey+10),b,25)
            return
        self.text((x+(12 if nx>=x else -48),y+(12 if ny>y else -36)),a,25)
        self.text((ex+(12 if ex==px else -67 if ex>px else 12),ey+(-39 if ey>py else 12 if ey<py else -36)),b,25)
        if label:self.text(label[:2],label[2],27,True)
    def title(self,t):self.text((800,25),t,40,True)
    def save(self):
        OUT.mkdir(parents=True,exist_ok=True); p=OUT/(self.name+'.png'); self.im.save(p,dpi=(300,300))
        (OUT/(self.name+'.svg')).write_text('\n'.join(self.svg+['</svg>']),encoding='utf-8'); return p

def modules():
    c=Canvas('function-modules',1020); c.title('叙事工坊功能模块划分'); c.box((600,110,1000,200),'叙事工坊',BLUE,37)
    xs=[50,440,830,1220]; centers=[x+165 for x in xs]; c.line([(800,200),(800,260)]); c.line([(215,260),(1385,260)])
    names=['账户与项目','内容建模','剧情编排','模拟与质量']
    leaves=[['注册与登录','项目维护','成员与权限'],['角色档案','世界设定','状态变量'],['剧情节点','选择分支','条件与效果'],['剧情模拟','结构检测','反馈处理']]
    for x,mid,name,items in zip(xs,centers,names,leaves):
        c.line([(mid,260),(mid,330)]); c.box((x,330,x+330,420),name,GOLD)
        c.line([(mid,420),(mid,455),(x+14,455),(x+14,770)])
        for j,label in enumerate(items):
            y=500+j*115; c.line([(x+14,y+40),(x+55,y+40)]); c.box((x+55,y,x+330,y+80),label,'white',32)
    c.text((800,925),'层级关系：系统 → 业务模块 → 子功能',30,True); return c.save()

def resources():
    c=Canvas('feature-resources',630); c.title('剧情编辑功能的 MVC 资源组织')
    rows=[('View','WorkspaceView.vue · 剧情图画布 · 节点与选项表单',BLUE),('Controller','StoryGraphController · StoryGraphDtos',GOLD),('Service','StoryGraphService · StoryGraphServiceImpl · ProjectAccessService',GREEN),('Mapper','StoryNodeMapper · StoryChoiceMapper',PURPLE),('Entity','StoryNode · StoryChoice',BLUE)]
    for i,(kind,value,color) in enumerate(rows):
        y=110+i*98; c.box((60,y,300,y+76),kind,color,31); c.line([(300,y+38),(390,y+38)],arrow=True); c.box((390,y,1540,y+76),value,'white',28)
    return c.save()

def classes():
    c=Canvas('feature-classes',1500); c.title('剧情编辑功能的核心类关系')
    c.entity(580,110,'StoryGraphController',['+ createNode(...)','+ createChoice(...)'],GOLD)
    c.entity(580,420,'StoryGraphService',['«interface»','+ createNode(...)','+ createChoice(...)'],GREEN)
    c.entity(580,740,'StoryGraphServiceImpl',['- accessService','- nodeMapper / choiceMapper'],GREEN)
    c.entity(40,740,'ProjectAccessService',['+ requireEditor(...)'],GREEN)
    c.entity(40,1070,'StoryNodeMapper',['+ insert(...) / selectById(...)'],PURPLE)
    c.entity(1120,1070,'StoryChoiceMapper',['+ insert(...) / selectList(...)'],PURPLE)
    c.entity(40,1330,'StoryNode',['id / projectId / nodeKey'],BLUE)
    c.entity(1120,1330,'StoryChoice',['sourceNodeId / targetNodeId'],BLUE)
    c.line([(800,279),(800,420)],True,True); c.text((820,333),'调用')
    c.line([(800,740),(800,632)],True,hollow=True); c.text((820,665),'实现')
    c.line([(580,820),(480,820)],True,True); c.text((530,770),'鉴权',27,True)
    for x,start in [(260,680),(1340,920)]:
        c.line([(start,909),(start,990),(x,990),(x,1070)],True,True)
        c.line([(x,1196),(x,1330)],True,True); c.text((x+20,1250),'持久化',28)
    return c.save()

def sequence():
    c=Canvas('feature-sequence',1340); c.title('新增选择分支的 UML 时序图'); xs=[150,475,800,1125,1450]
    labels=['页面\nWorkspaceView','控制器\nStoryGraphController','业务服务\nStoryGraphServiceImpl','权限服务\nProjectAccessService','数据访问\nNode / ChoiceMapper']
    for x,label in zip(xs,labels):
        c.box((x-140,110,x+140,225),label,BLUE,24); c.line([(x,225),(x,1240)],True)
    msgs=[(0,1,'1 提交分支 DTO',False),(1,2,'2 createChoice(...)',False),(2,3,'3 requireEditor(...)',False),(3,2,'4 通过权限检查',True),(2,4,'5 查询源节点、目标节点',False),(4,2,'6 返回节点数据',True),(2,2,'7 业务约束校验',False),(2,4,'8 insert(StoryChoice)',False),(4,2,'9 返回持久化结果',True),(2,1,'10 返回 ChoiceSummary',True),(1,0,'11 返回响应，更新画布',True)]
    for i,(a,b,label,dash) in enumerate(msgs):
        y=310+i*83
        if a==b:
            c.text((xs[a]+30,y-39),label,27); c.line([(xs[a],y),(xs[a]+75,y),(xs[a]+75,y+30),(xs[a],y+30)],arrow=True)
        else:
            c.text((min(xs[a],xs[b])+15,y-40),label,26); c.line([(xs[a],y),(xs[b],y)],dash,True)
    c.text((800,1280),'实线：调用；虚线：返回。校验失败时返回业务错误，不写入分支。',29,True); return c.save()

def accounts():
    c=Canvas('database-a-accounts',1120); c.title('（a）账户与项目')
    c.entity(60,210,'sys_user',['PK id','UK username']); c.entity(1100,210,'narrative_project',['PK id','FK owner_id'])
    c.entity(580,210,'project_member',['PK id','FK user_id','FK project_id','UK (project_id, user_id)'])
    c.entity(60,740,'world_entry',['PK id','FK project_id','entry_type / title'],GREEN)
    c.entity(1100,740,'detected_issue',['PK id','FK project_id','issue_type / target_id'],GOLD)
    c.line([(280,210),(280,130),(1320,130),(1320,210)])
    c.text((300,170),'1',25); c.text((1340,170),'0..*',25); c.text((800,87),'创建者 owner_id',27,True)
    c.rel([(500,290),(580,290)]); c.rel([(1100,290),(1020,290)])
    c.rel([(1400,379),(1400,740)])
    c.rel([(1180,379),(1180,620),(280,620),(280,740)])
    c.text((800,995),'PK：主键；FK：外键；UK：唯一约束。',30,True)
    c.text((800,1042),'1、0..1、0..* 分别表示一、零或一、零至多。',29,True); return c.save()

def characters():
    c=Canvas('database-b-characters',1290); c.title('（b）角色与知识')
    c.entity(60,140,'character_profile',['PK id','FK project_id','name'],GREEN)
    c.entity(1100,140,'story_node',['PK id','FK project_id','UK (project_id, node_key)'],GOLD)
    c.entity(580,140,'node_character',['PK, FK node_id','PK, FK character_id'],GREEN)
    c.entity(60,690,'character_relation',['PK id','FK project_id','FK source_character_id','FK target_character_id'],GREEN)
    c.entity(1100,690,'character_knowledge',['PK id','FK project_id','FK character_id','FK? acquired_node_id'],GREEN)
    c.rel([(500,240),(580,240)])
    c.rel([(1100,240),(1020,240)])
    c.rel([(180,352),(180,690)],label=(120,475,'源角色'))
    c.rel([(360,352),(360,690)],label=(420,553,'目标角色'))
    c.rel([(500,310),(550,310),(550,590),(1200,590),(1200,690)],label=(1020,550,'所属角色'))
    c.rel([(1400,352),(1400,690)],a='0..1',label=(1460,500,'获知节点'))
    c.text((800,1060),'node_character 的两个外键共同构成复合主键。',30,True)
    c.text((800,1110),'FK? 表示可空外键；跨子图出现的同名实体是同一张表。',29,True)
    c.text((800,1170),'各 project_id 均引用 narrative_project.id（项目 1 — 本表 0..*）。',29,True); return c.save()

def story():
    c=Canvas('database-c-story',1300); c.title('（c）剧情与规则')
    c.entity(60,140,'story_node',['PK id','FK project_id','UK (project_id, node_key)'],GOLD)
    c.entity(1100,140,'story_choice',['PK id','FK project_id','FK source_node_id','FK target_node_id'],GOLD)
    c.entity(60,780,'choice_condition',['PK id','FK choice_id','FK variable_id','operator / expected_value'],GREEN)
    c.entity(580,780,'state_variable',['PK id','FK project_id','value_type / initial_value'],GREEN)
    c.entity(1100,780,'state_effect',['PK id','FK choice_id','FK variable_id','operation / operand_value'],GREEN)
    c.rel([(500,215),(1100,215)],label=(800,170,'source_node_id · 起始节点'))
    c.rel([(500,310),(1100,310)],label=(800,265,'target_node_id · 目标节点'))
    c.rel([(1220,395),(1220,540),(280,540),(280,780)],label=(690,493,'选择包含条件'))
    c.rel([(1420,395),(1420,780)],label=(1240,625,'选择包含效果'))
    c.rel([(580,895),(500,895)]); c.rel([(1020,895),(1100,895)])
    c.text((800,1160),'各 project_id 均引用 narrative_project.id（项目 1 — 本表 0..*）。',29,True)
    c.text((800,1210),'条件和效果通过 variable_id 引用状态变量。',29,True); return c.save()

def testing():
    c=Canvas('database-d-testing',1450); c.title('（d）测试与反馈')
    c.entity(60,140,'sys_user',['PK id']); c.entity(580,140,'playtest_session',['PK id','FK project_id','FK tester_id','FK current_node_id'],PURPLE)
    c.entity(1100,140,'story_node',['PK id'],GOLD)
    c.entity(60,770,'test_feedback',['PK id','FK project_id','FK reporter_id','FK? session_id'],PURPLE)
    c.entity(580,770,'playtest_step',['PK id','FK session_id','FK node_id','FK? choice_id'],PURPLE)
    c.entity(1100,1110,'story_choice',['PK id'],GOLD)
    c.rel([(500,210),(580,210)]); c.rel([(1100,210),(1020,210)])
    c.rel([(180,266),(180,770)],label=(300,480,'提交反馈'))
    c.rel([(680,395),(680,570),(400,570),(400,770)],a='0..1',label=(505,525,'关联会话'))
    c.rel([(850,395),(850,770)],label=(910,540,'会话步骤'))
    c.rel([(1330,266),(1330,650),(1000,650),(1000,770)],label=(1195,607,'访问节点'))
    c.rel([(1100,1180),(850,1180),(850,1025)],a='0..1',label=(980,1220,'所选分支'))
    c.text((800,1320),'各 project_id 均引用 narrative_project.id（项目 1 — 本表 0..*）。',29,True)
    c.text((800,1370),'FK? 表示可空外键；跨子图出现的同名实体是同一张表。',29,True); return c.save()

def picture(d,path,title):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(4); p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
    s=p.add_run().add_picture(str(path),width=Inches(6.15)); s._inline.docPr.set('descr',title)
    return p

def main():
    QA.mkdir(parents=True,exist_ok=True); baseline=QA/'before.docx'
    if not baseline.exists():shutil.copy2(DOC,baseline)
    d=Document(baseline); shapes=list(d.inline_shapes)
    paths=[modules(),resources(),classes(),sequence(),accounts(),characters(),story(),testing()]
    for idx,path in [(28,paths[0]),(30,paths[3])]:
        s=shapes[idx]; rid=s._inline.graphic.graphicData.pic.blipFill.blip.embed; d.part.related_parts[rid]._blob=path.read_bytes()
        with Image.open(path) as im:ratio=im.height/im.width
        s.width=Inches(6.15); s.height=Inches(6.15*ratio)
        Paragraph(s._inline.getparent().getparent().getparent(),d._body).paragraph_format.keep_with_next=True
    old=shapes[29]._inline.getparent().getparent().getparent()
    for path,title in [(paths[1],'剧情编辑功能 MVC 资源组织'),(paths[2],'剧情编辑功能核心类关系')]:old.addprevious(picture(d,path,title)._p)
    old.getparent().remove(old)
    old=shapes[33]._inline.getparent().getparent().getparent(); original_caption=old.getnext()
    titles=['账户与项目','角色与知识','剧情与规则','测试与反馈']
    for n,path in enumerate(paths[4:]):
        old.addprevious(picture(d,path,titles[n])._p)
        if n==0:
            cap=deepcopy(original_caption); ts=list(cap.iter(qn('w:t')))
            ts[0].text='数据库 UML 逻辑结构图（a）账户与项目'
            for t in ts[1:]:t.text=''
            cp=Paragraph(cap,d._body); cp.paragraph_format.keep_with_next=False; cp.paragraph_format.keep_together=True
            old.addprevious(cap)
        else:
            cap=d.add_paragraph(f'图34（{"abcd"[n]}）数据库 UML 逻辑结构图：{titles[n]}'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.first_line_indent=Pt(0); cap.paragraph_format.space_after=Pt(8)
            cap.paragraph_format.keep_with_next=False; cap.paragraph_format.keep_together=True
            cap.paragraph_format.line_spacing=1.0
            for r in cap.runs:r.font.size=Pt(10.5)
            old.addprevious(cap._p)
    original_caption.getparent().remove(original_caption); old.getparent().remove(old)
    for p in d.paragraphs:
        if p.text.startswith('主事件流为：创作者填写节点和分支后提交'):
            p.text='图中以新增选择分支为例：创作者提交分支后，Controller 调用 createChoice；Service 首先检查编辑权限，再查询源节点和目标节点，校验节点属于当前项目且源节点不是结局节点，随后在事务中持久化 StoryChoice 并返回 ChoiceSummary。若权限或业务校验失败，统一异常处理器返回业务错误，页面保留输入。新增节点通过独立的 createNode 请求完成，校验节点标识与起点约束后保存。'
        if p.text=='保存剧情节点与选择分支的 UML 序列图':p.text='新增选择分支的 UML 时序图'
        if p.text.startswith('逻辑结构来源于需求分析中的业务实体。'):p.add_run(' 图34按业务域分为四个子图，同名实体表示同一张表；project_id 的统一归属关系在子图下方说明，以避免重复连线。')
    target=DOC
    try:d.save(target)
    except PermissionError:
        target=DOC.with_name('需求分析_图示重排.docx'); d.save(target)
    assert len(d.tables)==len(Document(baseline).tables)
    print('Output:',target)
    print('Updated; inline figures:',len(d.inline_shapes),'tables:',len(d.tables))

if __name__=='__main__':main()
